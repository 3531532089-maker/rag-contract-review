# 文档解析服务
# PDF 用 PyMuPDF (fitz) 提取文本，Word (.docx) 用 python-docx 提取文本。
# 扫描件（图片型 PDF）或图片型 Word 自动走 OCR（RapidOCR）。

import logging
import tempfile
from pathlib import Path

import fitz  # PyMuPDF 的导入名

from app.services.ocr_service import ocr_image_path

logger = logging.getLogger("contract-review.parser")

OCR_DPI = 200  # 扫描页渲染分辨率（dpi），越高越清晰但越慢


def _render_page_to_png(page) -> str:
    """把 PDF 页渲染成临时 PNG，返回临时文件路径。"""
    pix = page.get_pixmap(dpi=OCR_DPI)
    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    tmp.close()
    pix.save(tmp.name)
    return tmp.name


def _remove(path):
    try:
        Path(path).unlink(missing_ok=True)
    except OSError:
        pass


def extract_text_from_pdf(file_path: str) -> str:
    """
    从 PDF 中提取文本；无文字层的页（扫描件）自动 OCR。

    异常:
        ValueError: 加密 PDF
    """
    doc = fitz.open(file_path)

    if doc.is_encrypted:
        doc.close()
        raise ValueError(
            "该 PDF 文件已加密（设有密码保护）。请先使用 Adobe Acrobat 或在线工具"
            "去除密码保护后重新上传。本系统暂不支持解析加密的 PDF。"
        )

    text_parts = []
    for page_num, page in enumerate(doc, start=1):
        page_text = page.get_text()
        if page_text.strip():
            text_parts.append(f"--- 第{page_num}页 ---\n{page_text}")
            continue
        # 无文字层：可能是扫描件，尝试 OCR
        ocr_text = _ocr_page(page)
        if ocr_text.strip():
            text_parts.append(f"--- 第{page_num}页 ---\n{ocr_text}")

    doc.close()
    return "\n\n".join(text_parts)


def _ocr_page(page) -> str:
    """对单个无文字层的 PDF 页做 OCR。"""
    tmp = _render_page_to_png(page)
    try:
        return ocr_image_path(tmp)
    finally:
        _remove(tmp)


def _docx_images(file_path: str) -> list:
    """提取 Word 文档正文中内嵌的图片字节。"""
    from docx import Document

    doc = Document(file_path)
    images = []
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            try:
                images.append(rel.target_part.blob)
            except Exception:
                continue
    return images


def extract_text_from_docx(file_path: str) -> str:
    """
    从 Word (.docx) 中提取文本（含表格）。
    若正文没有任何文字，则对内嵌图片做 OCR（应对"扫描件塞进 Word"的情况）。
    """
    from docx import Document

    doc = Document(file_path)
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    text = "\n\n".join(parts)
    if text.strip():
        return text

    # 无文字层：对内嵌图片 OCR
    for blob in _docx_images(file_path):
        tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
        tmp.close()
        try:
            with open(tmp.name, "wb") as f:
                f.write(blob)
            ocr_text = ocr_image_path(tmp.name)
            if ocr_text.strip():
                parts.append(ocr_text)
        finally:
            _remove(tmp.name)

    return "\n\n".join(parts)


def parse_contract_file(file_path: str) -> dict:
    """
    自动识别文件类型（PDF 或 Word），提取文本（必要时 OCR）。

    返回:
        {"text": "...", "page_count": N, "file_type": "pdf"}

    异常:
        ValueError: 加密 PDF / 无法识别出任何文字
    """
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        doc = fitz.open(file_path)
        if doc.is_encrypted:
            doc.close()
            raise ValueError(
                "该 PDF 文件已加密（设有密码保护）。请先去除密码保护后重新上传。"
            )
        page_count = doc.page_count
        doc.close()

        text = extract_text_from_pdf(file_path)
        if not text.strip():
            raise ValueError(
                "PDF 文件中未检测到文本内容，OCR 也未能识别。"
                "可能是图片过于模糊，请上传清晰的扫描件或带文字层的 PDF。"
            )
        return {"text": text, "page_count": page_count, "file_type": "pdf"}

    elif suffix == ".docx":
        text = extract_text_from_docx(file_path)
        if not text.strip():
            raise ValueError(
                "Word 文件中未检测到文本内容，OCR 也未能识别。"
                "可能是纯图片型 Word（扫描件图片），请上传带文字的 Word 文档。"
            )
        return {"text": text, "page_count": None, "file_type": "word"}

    else:
        raise ValueError(f"不支持的文件格式: {suffix}。请上传 PDF 或 Word(.docx) 文件。")
